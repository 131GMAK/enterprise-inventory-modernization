#!/usr/bin/env python3
"""
Extract per-item product photos from an Excel order sheet whose pictures are
pasted screenshots of supplier tables (photo cell + text cells + black gridlines).

HOW IT WORKS (one pipeline for every image, no "strategies"):

  1. Read the drawing layer straight from the .xlsx XML (not via openpyxl's
     ws._images) so we keep what openpyxl throws away: the srcRect crop, and the
     fact that one media file can be placed several times with different crops.
  2. Convert every picture's anchor into a vertical span in points on the sheet,
     using the real row heights. Every row the picture covers is a "band".
  3. Map bands -> pixel boundaries inside the (cropped) picture. This is the PRIOR:
     accurate to a few pixels because Excel stretched the picture to the anchor box.
  4. Detect thin dark gridlines with a ridge detector and SNAP each predicted
     boundary to the nearest real line. Lines far from a predicted boundary are
     ignored, so dark photo content can never create a fake cut.
  5. Detect vertical gridlines, pick the photo column (the cell with the most
     "ink"), and save one file per item row.

KNOWN LIMIT: if the supplier's screenshot has a photo that straddles two table rows
(a floating picture, not confined to its cell), the cut still follows the row line, so
that photo is split between the two rows. That is a property of the source image.

Rows are matched to pictures by GEOMETRY (which rows the picture sits on), never by
"i-th picture == i-th item", so spacer rows, header rows, freight rows or a
missing picture cannot shift the matching.

Usage:
    pip install openpyxl numpy pillow
    python extract_images.py SAMPLE.xlsx -o images --debug-dir debug
    python extract_images.py SAMPLE.xlsx -o images --mode cells   # every cell
    python extract_images.py SAMPLE.xlsx -o images --mode raw     # every embedded picture, unsliced

Modes:
    photo (default) - one file per item ROW, cropped to just the photo cell.
                       This is already "separate files": one PNG per product.
    row             - one file per item row, the whole row (all columns).
    cells           - one file per item row PER CELL (photo + every text cell).
    raw             - no row/column slicing at all: one file per embedded
                       picture OBJECT in the drawing layer, cropped only by
                       whatever srcRect Excel applies. Use this to sanity-check
                       what's actually embedded, or when a picture isn't a
                       stitched table at all (e.g. a single product photo).
"""
from __future__ import annotations

import argparse
import csv
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import numpy as np
import openpyxl
from PIL import Image, ImageDraw

# --------------------------------------------------------------------------- #
# Tunables
# --------------------------------------------------------------------------- #
EMU_PER_PT = 12700          # 1 point = 12700 EMU (Excel's drawing unit)
MIN_ROW_OVERLAP = 0.5       # a picture "covers" a row if it spans >= 50% of it

RIDGE_K = 4                 # a line pixel must be darker than pixels k px on BOTH sides
RIDGE_DELTA = 35            # ...by at least this many grey levels
H_LINE_MIN = 0.45           # share of a pixel row that must be ridge to count as a line
V_LINE_MIN = 0.50           # same for pixel columns
SNAP_FRAC = 0.25            # snap tolerance = 25% of the smaller neighbouring row height
SNAP_MIN_PX = 6             # ...but never less than this
EDGE_PX = 6                 # vertical lines this close to the border are the outer frame
MIN_CELL_PX = 12            # ignore cells narrower/shorter than this
CELL_MARGIN = 2             # shave this many px off each side to drop the gridline itself
INK_LEVEL = 235             # a pixel darker than this counts as "ink" (not white paper)
LOW_CONFIDENCE = 1.5        # best/second-best ink ratio below this -> warn

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
R_ID = f"{{{NS['r']}}}id"
R_EMBED = f"{{{NS['r']}}}embed"


# --------------------------------------------------------------------------- #
# 1. Reading the drawing layer from the raw package
# --------------------------------------------------------------------------- #
@dataclass
class Picture:
    media: str                                   # path inside the zip
    crop: tuple[float, float, float, float]      # l, t, r, b as fractions
    col0: int                                    # 0-based first column
    y0: float                                    # top edge, points from sheet top
    y1: float                                    # bottom edge, points


def _resolve(base_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """{rId: (relationship kind, resolved target path)} for a package part."""
    path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if path not in z.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(z.read(path)):
        if rel.get("TargetMode") == "External":
            continue
        out[rel.get("Id")] = (rel.get("Type").rsplit("/", 1)[-1], _resolve(part, rel.get("Target")))
    return out


class RowGeometry:
    """Cumulative row tops in points, honouring custom, default and hidden rows."""

    def __init__(self, ws, extra_rows: int = 1000):
        default = ws.sheet_format.defaultRowHeight or 15.0
        n = ws.max_row + extra_rows
        self.height = [0.0] * (n + 2)            # 1-based
        self.top = [0.0] * (n + 3)
        for r in range(1, n + 1):
            dim = ws.row_dimensions.get(r)
            if dim is not None and dim.hidden:
                h = 0.0
            elif dim is not None and dim.height is not None:
                h = float(dim.height)
            else:
                h = float(default)
            self.height[r] = h
            self.top[r + 1] = self.top[r] + h

    def y(self, row0: int, off_emu: int) -> float:
        """Absolute y (pt) of an anchor point: 0-based row + EMU offset inside it.
        The offset is deliberately NOT clamped to the row height: if a file carries
        an overflowing offset it correctly spills into the following rows."""
        return self.top[row0 + 1] + off_emu / EMU_PER_PT


def _anchor_point(el, geom: RowGeometry) -> tuple[int, float]:
    col = int(el.findtext("xdr:col", namespaces=NS))
    row = int(el.findtext("xdr:row", namespaces=NS))
    row_off = int(el.findtext("xdr:rowOff", namespaces=NS))
    return col, geom.y(row, row_off)


def read_pictures(xlsx: Path, sheet_name: str, geom: RowGeometry) -> list[Picture]:
    pics: list[Picture] = []
    with zipfile.ZipFile(xlsx) as z:
        wb_root = ET.fromstring(z.read("xl/workbook.xml"))
        rid = next(s.get(R_ID) for s in wb_root.find("m:sheets", NS) if s.get("name") == sheet_name)
        sheet_part = _rels(z, "xl/workbook.xml")[rid][1]

        for kind, drawing_part in _rels(z, sheet_part).values():
            if kind != "drawing":
                continue
            drels = _rels(z, drawing_part)
            for anchor in ET.fromstring(z.read(drawing_part)):
                pic = anchor.find("xdr:pic", NS)
                if pic is None:                              # shapes, charts, ... skip
                    continue
                blip = pic.find("xdr:blipFill/a:blip", NS)
                if blip is None or blip.get(R_EMBED) not in drels:
                    continue

                tag = anchor.tag.split("}")[1]
                if tag == "twoCellAnchor":
                    col0, y0 = _anchor_point(anchor.find("xdr:from", NS), geom)
                    _, y1 = _anchor_point(anchor.find("xdr:to", NS), geom)
                elif tag == "oneCellAnchor":
                    col0, y0 = _anchor_point(anchor.find("xdr:from", NS), geom)
                    y1 = y0 + int(anchor.find("xdr:ext", NS).get("cy")) / EMU_PER_PT
                elif tag == "absoluteAnchor":
                    col0 = 0
                    y0 = int(anchor.find("xdr:pos", NS).get("y")) / EMU_PER_PT
                    y1 = y0 + int(anchor.find("xdr:ext", NS).get("cy")) / EMU_PER_PT
                else:
                    continue

                sr = pic.find("xdr:blipFill/a:srcRect", NS)
                crop = tuple(int(sr.get(k, 0)) / 100000 for k in "ltrb") if sr is not None else (0, 0, 0, 0)
                pics.append(Picture(drels[blip.get(R_EMBED)][1], crop, col0, y0, y1))
    return sorted(pics, key=lambda p: (p.y0, p.col0))


def load_image(xlsx: Path, pic: Picture) -> Image.Image:
    with zipfile.ZipFile(xlsx) as z:
        img = Image.open(io.BytesIO(z.read(pic.media))).convert("RGB")
    l, t, r, b = pic.crop                     # apply the crop Excel applies at display time
    w, h = img.size
    return img.crop((round(l * w), round(t * h), round(w * (1 - r)), round(h * (1 - b))))


# --------------------------------------------------------------------------- #
# 2. Reading the sheet: which rows are real items, and which batch are they in
# --------------------------------------------------------------------------- #
CONTAINER_RE = re.compile(r"^[A-Z]{4}\d{7}$")                  # e.g. TCNU5616606
BATCH_RE = re.compile(r"\b(ORDER|BATCH)\b", re.I)              # \b: "BORDER" won't match
NON_ITEM_RE = re.compile(r"\b(CTNS|FREIGHT|TOTAL|BALANCE|REMAIN|SEND TO ME)\b", re.I)


@dataclass
class RowInfo:
    batch: str
    sku: str
    is_item: bool


def _slug(text: str) -> str:
    return re.sub(r"\s+", "_", str(text).strip()).upper()      # collapses double spaces too


def _safe(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text)               # filesystem-safe


def scan_sheet(ws, default_batch: str | None, sku_col=4, batch_col=2, last_col=13):
    initial = ws.cell(1, batch_col).value
    batch = default_batch or (_slug(initial) if initial else "BATCH_DEFAULT")
    rows: dict[int, RowInfo] = {}
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, batch_col).value
        if b and BATCH_RE.search(str(b)):
            batch = _slug(b)                                   # batch header row
            rows[r] = RowInfo(batch, "", False)
            continue
        raw = ws.cell(r, sku_col).value
        sku = str(raw).strip() if raw is not None else ""
        row_text = " ".join(str(ws.cell(r, c).value) for c in range(1, last_col + 1)
                            if ws.cell(r, c).value is not None)
        is_item = (bool(sku) and sku.upper() != "ITEM NO"
                   and not CONTAINER_RE.match(sku) and not NON_ITEM_RE.search(row_text))
        rows[r] = RowInfo(batch, sku, is_item)
    return rows


def table_last_col(ws, sku_col=4) -> int:
    """1-based last column of the order table = end of the contiguous header block
    that contains the ITEM NO header. Pictures starting to the right of it are
    stray (e.g. a duplicate pasted into a scratch area) and are skipped."""
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, sku_col).value).strip().upper() == "ITEM NO":
            c = sku_col
            while ws.cell(r, c + 1).value not in (None, ""):
                c += 1
            return c
    return 13


# --------------------------------------------------------------------------- #
# 3. Image analysis
# --------------------------------------------------------------------------- #
def ridge_profile(gray: np.ndarray, axis: int) -> np.ndarray:
    """For each pixel row (axis=0) or column (axis=1): the fraction of pixels that
    are a thin DARK RIDGE, i.e. darker than the pixels k px away on both sides.

    Why not 'fraction of dark pixels'? Screenshot gridlines are often grey or
    anti-aliased (missed by a fixed threshold), while dark photo regions are wide
    (false positives). A ridge is thin, so it separates the two."""
    g = gray.astype(np.int16)
    if axis == 1:
        g = g.T
    n, k = g.shape[0], RIDGE_K
    score = np.zeros(n)
    if n > 2 * k:
        core = g[k:n - k]
        ridge = (core < g[:n - 2 * k] - RIDGE_DELTA) & (core < g[2 * k:] - RIDGE_DELTA)
        score[k:n - k] = ridge.mean(axis=1)
    return score


def find_lines(score: np.ndarray, threshold: float) -> list[float]:
    groups: list[list[int]] = []
    for i in np.flatnonzero(score >= threshold):
        if groups and i - groups[-1][-1] <= 2:
            groups[-1].append(int(i))
        else:
            groups.append([int(i)])
    return [float(np.mean(g)) for g in groups]


def predict_boundaries(rows: list[int], pic: Picture, geom: RowGeometry, img_h: int) -> list[float]:
    """Pixel y of every row boundary, from sheet geometry alone."""
    span = pic.y1 - pic.y0
    edges = [geom.top[r] for r in rows] + [geom.top[rows[-1] + 1]]
    px = [(e - pic.y0) / span * img_h for e in edges]
    px[0], px[-1] = max(0.0, px[0]), min(float(img_h), px[-1])
    return px


def snap_boundaries(pred: list[float], h_lines: list[float]) -> tuple[list[float], list[str]]:
    """Snap each INTERNAL predicted boundary to the nearest detected gridline within
    tolerance; otherwise keep the prediction. Outer edges stay where predicted."""
    final, how = list(pred), ["edge"]
    for i in range(1, len(pred) - 1):
        gap = min(pred[i] - pred[i - 1], pred[i + 1] - pred[i])
        tol = max(SNAP_MIN_PX, SNAP_FRAC * gap)
        near = [y for y in h_lines if abs(y - pred[i]) <= tol]
        if near:
            final[i] = min(near, key=lambda y: abs(y - pred[i]))
            how.append("snapped")
        else:
            how.append("predicted")
    how.append("edge")
    for i in range(1, len(final)):                       # keep boundaries strictly increasing
        if final[i] <= final[i - 1] + 1:
            final[i], how[i] = pred[i], "predicted"
    return final, how


def column_edges(gray: np.ndarray) -> list[float]:
    width = gray.shape[1]
    inner = [x for x in find_lines(ridge_profile(gray, 1), V_LINE_MIN) if EDGE_PX < x < width - EDGE_PX]
    edges = [0.0] + inner + [float(width)]
    keep = [edges[0]]
    for x in edges[1:]:
        if x - keep[-1] >= MIN_CELL_PX:
            keep.append(x)
    keep[-1] = float(width)
    return keep


def pick_photo_column(gray: np.ndarray, x_edges: list[float]) -> tuple[int, float]:
    """Photo cell = the column with the most non-white pixels. Text cells are mostly
    white paper; a photo fills its cell. Returns (index, best/second-best ratio)."""
    ink = []
    for x0, x1 in zip(x_edges[:-1], x_edges[1:]):
        cell = gray[:, int(x0) + CELL_MARGIN:int(x1) - CELL_MARGIN]
        ink.append(float((cell < INK_LEVEL).mean()) if cell.size else 0.0)
    order = np.argsort(ink)[::-1]
    best = int(order[0])
    ratio = ink[best] / max(ink[int(order[1])], 1e-6) if len(order) > 1 else float("inf")
    return best, ratio


def trim_dark_edges(gray: np.ndarray, box, max_px=6, dark=110, frac=0.8):
    """Shave leftover gridline slivers: while an outermost pixel line of the crop is
    >= 80% dark, drop it (at most max_px per side). The ridge detector can't see
    lines within RIDGE_K px of the picture border, so the outer frame ends up here."""
    x0, y0, x1, y1 = box
    for _ in range(max_px):
        changed = False
        if (gray[y0:y1, x0] < dark).mean() >= frac: x0 += 1; changed = True
        if (gray[y0:y1, x1 - 1] < dark).mean() >= frac: x1 -= 1; changed = True
        if (gray[y0, x0:x1] < dark).mean() >= frac: y0 += 1; changed = True
        if (gray[y1 - 1, x0:x1] < dark).mean() >= frac: y1 -= 1; changed = True
        if not changed or x1 - x0 < MIN_CELL_PX or y1 - y0 < MIN_CELL_PX:
            break
    return x0, y0, x1, y1


# --------------------------------------------------------------------------- #
# 4. Orchestration
# --------------------------------------------------------------------------- #
def debug_overlay(img, pred, final, x_edges, path):
    im = img.copy()
    d = ImageDraw.Draw(im)
    for y in pred:
        d.line([(0, y), (im.width, y)], fill=(0, 90, 255), width=1)     # blue = predicted
    for y in final:
        d.line([(0, y), (im.width, y)], fill=(255, 0, 0), width=1)      # red  = final cut
    for x in x_edges:
        d.line([(x, 0), (x, im.height)], fill=(0, 170, 0), width=1)     # green = column cut
    im.save(path)


def extract(xlsx, out_dir, sheet=None, default_batch=None, mode="photo", debug_dir=None):
    xlsx, out_dir = Path(xlsx), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    if debug_dir:
        Path(debug_dir).mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    geom = RowGeometry(ws)
    info = scan_sheet(ws, default_batch)
    last_col = table_last_col(ws)
    pictures = read_pictures(xlsx, ws.title, geom)

    n_items = sum(1 for i in info.values() if i.is_item)
    print(f"Sheet '{ws.title}': {n_items} item rows, {len(pictures)} pictures, "
          f"table spans columns 1..{last_col}")

    if mode == "raw":
        # No row/column detection at all: every picture OBJECT becomes its own
        # file, in the order it appears top-to-bottom. Two anchors that reuse
        # the same media file (a picture pasted twice with different crops)
        # still produce two separate files here, because each is cropped by
        # its own srcRect.
        manifest = []
        for idx, pic in enumerate(pictures):
            img = load_image(xlsx, pic)
            name = f"{idx:02d}_{Path(pic.media).stem}.png"
            img.save(out_dir / name)
            manifest.append({"file": name, "picture": idx, "media": Path(pic.media).name,
                             "width": img.width, "height": img.height})
            print(f"  saved {name}  ({img.width}x{img.height})")
        print(f"\nSaved {len(manifest)} file(s) to {out_dir}")
        if manifest:
            with open(out_dir / "manifest.csv", "w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=list(manifest[0]))
                w.writeheader()
                w.writerows(manifest)
        return manifest

    manifest, claimed = [], {}
    for idx, pic in enumerate(pictures):
        tag = f"picture {idx} ({Path(pic.media).name})"
        if pic.col0 >= last_col:
            print(f"  SKIP  {tag}: starts in column {pic.col0 + 1}, outside the order table")
            continue

        # rows this picture covers = rows it overlaps by >= MIN_ROW_OVERLAP
        rows = []
        for r in range(1, len(geom.height) - 1):
            h = geom.height[r]
            if h > 0 and geom.top[r] < pic.y1 and geom.top[r + 1] > pic.y0:
                overlap = min(pic.y1, geom.top[r + 1]) - max(pic.y0, geom.top[r])
                if overlap / h >= MIN_ROW_OVERLAP:
                    rows.append(r)
        if not rows:
            print(f"  SKIP  {tag}: overlaps no rows")
            continue

        img = load_image(xlsx, pic)
        gray = np.array(img.convert("L"))
        pred = predict_boundaries(rows, pic, geom, img.height)
        final, how = snap_boundaries(pred, find_lines(ridge_profile(gray, 0), H_LINE_MIN))

        x_edges = column_edges(gray)
        n_cols = len(x_edges) - 1
        photo_col, ratio = pick_photo_column(gray, x_edges) if n_cols > 1 else (0, float("inf"))

        snapped = how.count("snapped")
        print(f"  OK    {tag}: rows {rows[0]}-{rows[-1]} ({len(rows)}), "
              f"{snapped}/{max(len(rows) - 1, 0)} cuts snapped to gridlines, "
              f"{n_cols} column(s), photo=col {photo_col + 1}"
              + ("  [LOW CONFIDENCE photo column]" if ratio < LOW_CONFIDENCE else ""))
        if debug_dir:
            debug_overlay(img, pred, final, x_edges, Path(debug_dir) / f"picture{idx:02d}.png")

        for i, r in enumerate(rows):
            ri = info.get(r)
            if ri is None or not ri.is_item:
                print(f"        row {r}: covered by picture but not an item row -> skipped")
                continue
            if r in claimed:
                print(f"        row {r}: already filled by picture {claimed[r]} -> skipped")
                continue
            claimed[r] = idx

            y0, y1 = final[i], final[i + 1]
            # legacy naming: the original script used the 0-based row index (excel row - 1)
            base = _safe(f"{ri.batch}_{r - 1}_{ri.sku}")

            cells = []
            if mode == "row":
                cells = [(0, 0.0, float(img.width))]
            elif mode == "cells":
                cells = [(c, x_edges[c], x_edges[c + 1]) for c in range(n_cols)]
            else:
                cells = [(photo_col, x_edges[photo_col], x_edges[photo_col + 1])]

            for c, x0, x1 in cells:
                m = 0 if mode == "row" else CELL_MARGIN
                box = (int(x0) + m, int(y0) + m, int(x1) - m, int(y1) - m)
                if box[2] - box[0] < MIN_CELL_PX or box[3] - box[1] < MIN_CELL_PX:
                    continue
                if mode != "row":
                    box = trim_dark_edges(gray, box)
                name = f"{base}.png" if mode != "cells" else f"{base}__c{c + 1}.png"
                img.crop(box).save(out_dir / name)
                manifest.append({"file": name, "sheet_row": r, "batch": ri.batch, "sku": ri.sku,
                                 "picture": idx, "media": Path(pic.media).name,
                                 "top_cut": how[i], "bottom_cut": how[i + 1],
                                 "photo_col_ratio": round(ratio, 2) if ratio != float("inf") else "",
                                 "status": "matched"})

    saved_count = len(manifest)                  # file count, before 'no image detected' rows are added below
    missing = [r for r, i in info.items() if i.is_item and r not in claimed]
    for r in sorted(missing):                     # one row per gap, so tests can assert on manifest.csv
        ri = info[r]                              # directly instead of parsing console output
        manifest.append({"file": "", "sheet_row": r, "batch": ri.batch, "sku": ri.sku,
                         "picture": "", "media": "", "top_cut": "", "bottom_cut": "",
                         "photo_col_ratio": "", "status": "no image detected"})
    manifest.sort(key=lambda m: m["sheet_row"])   # sheet order, matched and unmatched interleaved

    print(f"\nSaved {saved_count} file(s) for {len(claimed)} item rows to {out_dir}")
    print(f"{len(missing)} item rows have no image detected (e.g. rows {missing[:8]}...)" if missing else "")
    if manifest:
        with open(out_dir / "manifest.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(manifest[0]))
            w.writeheader()
            w.writerows(manifest)
    return manifest


if __name__ == "__main__":
    # Expected layout (adjust ROOT below if yours differs):
    #   aws-inventory-sync-engine/
    #     etl/extract_images.py   <- this file
    #     data/SAMPLE.xlsx        <- input manifests
    #     images/                 <- output (created if missing)
    HERE = Path(__file__).resolve().parent      # .../aws-inventory-sync-engine/etl
    ROOT = HERE.parent                          # .../aws-inventory-sync-engine
    DATA_DIR = ROOT / "data"
    DEFAULT_OUT = ROOT / "images"

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel", nargs="?", default=DATA_DIR / "SAMPLE.xlsx",
                    help="Excel file. A bare filename (e.g. SAMPLE.xlsx) is looked up in "
                         f"{DATA_DIR}; a path containing a folder (./other/file.xlsx, "
                         "an absolute path, ..\\elsewhere\\file.xlsx) is used exactly as given.")
    ap.add_argument("-o", "--out", default=DEFAULT_OUT,
                    help=f"output folder (default: {DEFAULT_OUT})")
    ap.add_argument("--sheet", help="sheet name (default: active sheet)")
    ap.add_argument("--batch", help="fallback batch code before the first batch header row")
    ap.add_argument("--mode", choices=["photo", "row", "cells", "raw"], default="photo",
                    help="photo = photo cell per item (default); row = whole row per item; "
                         "cells = every cell per item; raw = every embedded picture, unsliced")
    ap.add_argument("--debug-dir", help="write overlays showing predicted/final cuts for review")
    a = ap.parse_args()

    excel_path = Path(a.excel)
    if excel_path.parent == Path("."):           # bare filename, no folder given -> data/
        excel_path = DATA_DIR / excel_path.name
    if not excel_path.exists():
        raise SystemExit(f"Excel file not found: {excel_path}\n"
                          f"(bare filenames are looked up in {DATA_DIR})")

    extract(excel_path, a.out, a.sheet, a.batch, a.mode, a.debug_dir)
